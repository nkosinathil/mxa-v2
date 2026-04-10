import os, re, shutil, sqlite3, zipfile, hashlib, json, time
from pathlib import Path
from typing import List, Dict, Any
from .dashboard import generate_dashboard

def _rel_to_case(case_dir, path):
    value = str(path or '').strip()
    if not value:
        return ''
    try:
        return str(Path(value).resolve().relative_to(Path(case_dir).resolve())).replace('\\','/')
    except Exception:
        return os.path.basename(value)

from .db import DB
from .ocr import ocr_image, IMAGE_EXTS
from .categorizer import load_category_config
from .utils import extract_image_exif_metadata
from .image_metadata import extract_image_info, export_to_csv
from .parsers import WhatsAppHTMLParser, WhatsAppPDFParser, MessagesParser, CallsParser, EmailParser, AudioParser

MEDIA_EXTENSIONS = IMAGE_EXTS.union({'.pdf','.doc','.docx','.xls','.xlsx','.xlsm','.ppt','.pptx','.txt','.csv','.tsv','.json','.xml','.html','.htm','.md','.msg','.eml','.mp4','.mov','.avi','.mkv','.3gp','.webm','.mp3','.wav','.aac','.m4a','.ogg','.opus','.amr','.zip','.7z','.rar','.heics','.dng','.arw','.nef','.cr2'})

PARSERS = [WhatsAppHTMLParser(), WhatsAppPDFParser(), MessagesParser(), CallsParser(), EmailParser(), AudioParser()]

PARSER_MODE_MAP = {
    'whatsapp': (WhatsAppHTMLParser, WhatsAppPDFParser),
    'texts': (MessagesParser,),
    'calls': (CallsParser,),
    'emails': (EmailParser,),
    'audio': (AudioParser,),
}

def _normalize_selected_modes(selected_modes=None):
    defaults = {'whatsapp': True, 'texts': True, 'calls': True, 'emails': True, 'audio': True, 'photos': True, 'files': True}
    if not isinstance(selected_modes, dict):
        return defaults
    normalized = defaults.copy()
    for k, v in selected_modes.items():
        normalized[str(k).strip().lower()] = bool(v)
    return normalized

def _enabled_parsers(selected_modes=None):
    selected = _normalize_selected_modes(selected_modes)
    enabled_types = []
    for mode, parser_types in PARSER_MODE_MAP.items():
        if selected.get(mode, False):
            enabled_types.extend(parser_types)
    enabled = []
    seen = set()
    for parser in PARSERS:
        cls = parser.__class__
        if cls in enabled_types and cls not in seen:
            enabled.append(parser)
            seen.add(cls)
    return enabled

def _mode_from_parser(parser):
    cls = parser.__class__
    for mode, parser_types in PARSER_MODE_MAP.items():
        if cls in parser_types:
            return mode
    return 'unknown'


def safe_case_name(name):
    value = str(name or 'CASE-001').strip()
    value = re.sub(r'[^A-Za-z0-9._-]+', '_', value)
    return value or 'CASE-001'


def ensure_dirs(base_output, case_no):
    case_dir = os.path.join(base_output, safe_case_name(case_no))
    media_dir = os.path.join(case_dir, 'media')
    dashboard_dir = os.path.join(case_dir, 'dashboard')
    logs_dir = os.path.join(case_dir, 'logs')
    extracted_dir = os.path.join(case_dir, 'extracted')
    preview_dir = os.path.join(case_dir, 'previews')
    for p in [case_dir, media_dir, dashboard_dir, logs_dir, extracted_dir, preview_dir]:
        os.makedirs(p, exist_ok=True)
    return case_dir, media_dir, dashboard_dir, logs_dir, extracted_dir, preview_dir


def sha256_file(path):
    h = hashlib.sha256()
    with open(path,'rb') as f:
        for chunk in iter(lambda: f.read(1024*1024), b''):
            h.update(chunk)
    return h.hexdigest()


def discover_files(root):
    for current_root, _dirs, files in os.walk(root):
        for fname in files:
            yield os.path.join(current_root, fname)




def _safe_preview_name(path: str, suffix: str = '') -> str:
    base = os.path.basename(path or 'item')
    stem, ext = os.path.splitext(base)
    stem = re.sub(r'[^A-Za-z0-9._-]+', '_', stem)
    return f"{stem}{suffix}"


def _write_text(path: str, value: str):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, 'w', encoding='utf-8', errors='ignore') as f:
        f.write(value or '')


def extract_attachment_exif(image_path):
    try:
        exif = extract_image_exif_metadata(str(image_path)) or {}
        lat = exif.get('gps_lat')
        lon = exif.get('gps_lon')
        gps_ts = exif.get('exif_datetime_original') or exif.get('gps_timestamp') or ''
        return {
            'gps_lat': float(lat) if lat is not None else None,
            'gps_lon': float(lon) if lon is not None else None,
            'gps_alt': exif.get('gps_alt'),
            'gps_timestamp': gps_ts,
            'gps_source': exif.get('gps_source') or 'exif',
            'gps_confidence': exif.get('gps_confidence') if exif.get('gps_confidence') is not None else (1.0 if lat is not None and lon is not None else None),
            'exif_make': exif.get('exif_make') or '',
            'exif_model': exif.get('exif_model') or '',
            'exif_datetime_original': gps_ts,
            'has_gps': int(exif.get('has_gps') or (1 if lat is not None and lon is not None else 0)),
        }
    except Exception:
        return {
            'gps_lat': None,
            'gps_lon': None,
            'gps_alt': None,
            'gps_timestamp': '',
            'gps_source': 'exif',
            'gps_confidence': None,
            'exif_make': '',
            'exif_model': '',
            'exif_datetime_original': '',
            'has_gps': 0,
        }

def generate_preview_assets(media_paths, preview_dir, progress):
    image_map = {}
    doc_map = {}
    try:
        from PIL import Image
    except Exception:
        Image = None
    try:
        import pdfplumber
    except Exception:
        pdfplumber = None
    try:
        from pypdf import PdfReader
    except Exception:
        PdfReader = None
    try:
        from docx import Document as DocxDocument
    except Exception:
        DocxDocument = None
    try:
        import openpyxl
    except Exception:
        openpyxl = None

    media_map = {}

    img_dir = os.path.join(preview_dir, 'images')
    doc_dir = os.path.join(preview_dir, 'docs')
    media_dir = os.path.join(preview_dir, 'media')
    os.makedirs(img_dir, exist_ok=True)
    os.makedirs(doc_dir, exist_ok=True)
    os.makedirs(media_dir, exist_ok=True)

    for src in media_paths or []:
        if not src or not os.path.exists(src):
            continue
        ext = os.path.splitext(src)[1].lower()
        base = os.path.basename(src).lower()
        try:
            if ext in IMAGE_EXTS and Image is not None:
                out_name = _safe_preview_name(src, '.jpg')
                out_path = os.path.join(img_dir, out_name)
                if not os.path.exists(out_path):
                    im = Image.open(src)
                    try:
                        im = im.convert('RGB')
                    except Exception:
                        pass
                    im.thumbnail((640, 640))
                    im.save(out_path, format='JPEG', quality=82, optimize=True)
                image_map[base] = out_path
                continue
            if ext in ['.pdf','.doc','.docx','.xls','.xlsx','.xlsm','.csv','.txt','.log','.json','.xml','.html','.htm','.md']:
                out_name = _safe_preview_name(src, '.preview.txt')
                out_path = os.path.join(doc_dir, out_name)
                if not os.path.exists(out_path):
                    preview_text = ''
                    if ext == '.pdf':
                        chunks=[]
                        if pdfplumber is not None:
                            with pdfplumber.open(src) as pdf:
                                for page in pdf.pages[:8]:
                                    chunks.append(page.extract_text() or '')
                        elif PdfReader is not None:
                            reader = PdfReader(src)
                            for page in reader.pages[:8]:
                                chunks.append(page.extract_text() or '')
                        preview_text = '\n\n'.join([c for c in chunks if c]).strip()
                    elif ext in ['.docx'] and DocxDocument is not None:
                        doc = DocxDocument(src)
                        preview_text = '\n'.join([p.text for p in doc.paragraphs if p.text.strip()][:400])
                    elif ext in ['.xlsx','.xlsm'] and openpyxl is not None:
                        wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
                        sheet = wb[wb.sheetnames[0]]
                        lines = [f'Sheet: {sheet.title}']
                        for i,row in enumerate(sheet.iter_rows(values_only=True)):
                            if i >= 100: break
                            vals = ['' if v is None else str(v) for v in row[:16]]
                            lines.append('\t'.join(vals))
                        preview_text = '\n'.join(lines)
                    else:
                        with open(src, 'r', encoding='utf-8', errors='ignore') as f:
                            preview_text = f.read(50000)
                    _write_text(out_path, preview_text or f'No document preview available for: {os.path.basename(src)}')
                doc_map[base] = out_path
        except Exception as exc:
            if callable(progress):
                progress(f"[progress] Preview cache skipped for {src}: {exc}")
    return image_map, doc_map, media_map
def extract_zip_to(zip_path, target_root):
    base = os.path.splitext(os.path.basename(zip_path))[0]
    out_dir = os.path.join(target_root, base)
    os.makedirs(out_dir, exist_ok=True)
    try:
        with zipfile.ZipFile(zip_path, 'r') as zf:
            zf.extractall(out_dir)
        return out_dir
    except Exception:
        return None


def build_media_index(input_root, extracted_dir, media_dir, progress, logs_dir=None):
    hash_index = {}
    media_index = {}
    copied = []
    source_index = {}
    scan_roots = [input_root]
    folder_stats = {}

    for path in list(discover_files(input_root)):
        if path.lower().endswith('.zip'):
            progress(f"[progress] Extracting ZIP: {path}")
            extracted = extract_zip_to(path, extracted_dir)
            if extracted:
                scan_roots.append(extracted)

    seen_src = set()
    for root in scan_roots:
        for current_root, _dirs, files in os.walk(root, topdown=True):
            total_files = 0
            image_files = 0
            ext_counts = {}
            for fname in files:
                path = os.path.join(current_root, fname)
                norm = os.path.normpath(path)
                if norm in seen_src:
                    continue
                seen_src.add(norm)
                total_files += 1
                ext = os.path.splitext(path)[1].lower() or '[no_ext]'
                ext_counts[ext] = ext_counts.get(ext, 0) + 1
                if ext in IMAGE_EXTS:
                    image_files += 1
                if ext not in MEDIA_EXTENSIONS:
                    continue
                try:
                    file_hash = sha256_file(path)
                except Exception:
                    file_hash = None
                if file_hash and file_hash in hash_index:
                    target = hash_index[file_hash]
                else:
                    base = os.path.basename(path)
                    target = os.path.join(media_dir, base)
                    stem, suf = os.path.splitext(base)
                    counter = 2
                    while os.path.exists(target):
                        target = os.path.join(media_dir, f"{stem}__{counter}{suf}")
                        counter += 1
                    shutil.copy2(path, target)
                    if file_hash:
                        hash_index[file_hash] = target
                    copied.append(target)
                media_index.setdefault(os.path.basename(path).lower(), []).append(target)
                source_index[target] = path
            folder_stats[current_root] = {
                'total_files': total_files,
                'image_files': image_files,
                'extensions': dict(sorted(ext_counts.items(), key=lambda kv: kv[0]))
            }

    if logs_dir:
        try:
            os.makedirs(logs_dir, exist_ok=True)
            log_path = os.path.join(logs_dir, 'input_folder_processing_summary.log')
            with open(log_path, 'w', encoding='utf-8') as lf:
                for folder in sorted(folder_stats.keys()):
                    stats = folder_stats[folder]
                    lf.write(f"[SCAN] Folder: {folder}\n")
                    lf.write(f"  Total files: {stats['total_files']}\n")
                    lf.write(f"  Images: {stats['image_files']}\n")
                    media_count = sum(v for k,v in stats['extensions'].items() if k in MEDIA_EXTENSIONS or k in IMAGE_EXTS)
                    lf.write(f"  Supported media/docs: {media_count}\n")
                    lf.write("  Types:\n")
                    for ext, count in stats['extensions'].items():
                        lf.write(f"    {ext}: {count}\n")
                    lf.write("\n")
            progress(f"[progress] Input folder summary log written: {log_path}")
        except Exception as exc:
            progress(f"[progress] Input folder summary log skipped: {exc}")

    return media_index, copied, scan_roots, source_index
def normalize_record_id(row):
    payload = '||'.join([
        str(row.get('mode') or ''), str(row.get('source_file') or ''), str(row.get('timestamp') or ''),
        str(row.get('sender') or ''), str(row.get('receiver') or ''), str(row.get('subject') or ''),
        str(row.get('body') or row.get('message') or ''), str(row.get('attachment_name') or row.get('attachment') or '')
    ])
    return hashlib.sha256(payload.encode('utf-8', errors='ignore')).hexdigest()[:32]


def parse_with_registered_parsers(path, context, enabled_parsers=None, timing=None):
    parsers = enabled_parsers if enabled_parsers is not None else PARSERS
    for parser in parsers:
        mode_name = _mode_from_parser(parser)
        started = time.perf_counter()
        try:
            if parser.can_parse(path):
                rows = parser.parse(path, context) or []
                elapsed = time.perf_counter() - started
                if isinstance(timing, dict):
                    bucket = timing.setdefault(mode_name, {'files': 0, 'seconds': 0.0, 'rows': 0})
                    bucket['files'] += 1
                    bucket['seconds'] += elapsed
                    bucket['rows'] += len(rows)
                if rows:
                    return rows
        except Exception as exc:
            log = context.get('log')
            if callable(log):
                log(f"Parser failed on {path}: {exc}")
    return []


def collect_records(scan_roots, media_index, category_config, progress, case_dir, transcribe_audio=False, audio_max_transcription_seconds=None, preview_image_map=None, preview_doc_map=None, preview_media_map=None, source_index=None, selected_modes=None, parser_timing=None):
    records=[]
    attachment_rows=[]
    seen=set()
    selected_modes = _normalize_selected_modes(selected_modes)
    enabled_parsers = _enabled_parsers(selected_modes)
    for root in scan_roots:
        progress(f"[progress] Scanning folder: {root}")
        for path in discover_files(root):
            context = {'log': progress, 'owner_name': 'You', 'category_config': category_config, 'transcribe_audio': transcribe_audio, 'audio_max_transcription_seconds': audio_max_transcription_seconds}
            parsed = parse_with_registered_parsers(path, context, enabled_parsers=enabled_parsers, timing=parser_timing)
            if not parsed:
                continue
            progress(f"[progress] Parsed {len(parsed)} items from file: {path}")
            for row in parsed:
                row_mode = str(row.get('mode') or '').strip().lower()
                if row_mode.startswith('whatsapp') and not selected_modes.get('whatsapp', False):
                    continue
                if row_mode in ('texts', 'text messages', 'sms', 'messages') and not selected_modes.get('texts', False):
                    continue
                if row_mode in ('calls', 'call logs', 'phone calls', 'voice calls') and not selected_modes.get('calls', False):
                    continue
                if row_mode in ('emails', 'email') and not selected_modes.get('emails', False):
                    continue
                if row_mode in ('audio', 'audio files', 'voice notes') and not selected_modes.get('audio', False):
                    continue
                row.setdefault('source_file', path)
                row['id'] = normalize_record_id(row)
                key = row['id']
                if key in seen:
                    continue
                seen.add(key)
                att_name = (row.get('attachment_name') or row.get('attachment') or '').strip()
                if att_name:
                    candidates = media_index.get(os.path.basename(att_name).lower(), [])
                    media_path = candidates[0] if isinstance(candidates, list) and candidates else (candidates or '')
                    rel_media_path = _rel_to_case(case_dir, media_path) if media_path else ''
                    row['attachment_path'] = rel_media_path
                    row['media_path'] = rel_media_path
                    row['media_exists'] = bool(media_path and os.path.exists(media_path))
                    row['media_deleted'] = not row['media_exists']
                    if preview_image_map is not None:
                        pimg = preview_image_map.get(os.path.basename(att_name).lower(), '')
                        row['preview_image_path'] = _rel_to_case(case_dir, pimg) if pimg else ''
                    if preview_doc_map is not None:
                        pdoc = preview_doc_map.get(os.path.basename(att_name).lower(), '')
                        row['preview_doc_path'] = _rel_to_case(case_dir, pdoc) if pdoc else ''
                    if preview_media_map is not None:
                        pmed = preview_media_map.get(os.path.basename(att_name).lower(), '')
                        row['preview_media_path'] = _rel_to_case(case_dir, pmed) if pmed else ''
                    att = {
                        'id': hashlib.sha256(f"{row['id']}||{att_name}||{media_path}".encode()).hexdigest()[:32],
                        'communication_id': row['id'],
                        'mode': row.get('mode'),
                        'source_file': row.get('source_file'),
                        'attachment_name': os.path.basename(att_name),
                        'attachment': os.path.basename(att_name),
                        'attachment_path': rel_media_path,
                        'media_path': rel_media_path,
                        'found_status': 'found' if media_path and os.path.exists(media_path) else 'missing',
                        'preview_image_path': _rel_to_case(case_dir, preview_image_map.get(os.path.basename(att_name).lower(), '')) if preview_image_map else '',
                        'preview_doc_path': _rel_to_case(case_dir, preview_doc_map.get(os.path.basename(att_name).lower(), '')) if preview_doc_map else '',
                        'preview_media_path': _rel_to_case(case_dir, preview_media_map.get(os.path.basename(att_name).lower(), '')) if preview_media_map else '',
                        'attachment_type': os.path.splitext(att_name)[1].lower().lstrip('.'),
                        'file_ext': os.path.splitext(att_name)[1].lower(),
                        'ocr_status': '', 'ocr_text':'', 'text_detected':'no', 'reason':'', 'width':None, 'height':None,
                    }
                    if att['found_status'] == 'found' and att['file_ext'] in IMAGE_EXTS:
                        progress(f"[progress] OCR check: {media_path}")
                        o = ocr_image(media_path)
                        att.update(o)
                        att.update(extract_attachment_exif(media_path))
                        if o.get('ocr_text'):
                            row['ocr_text'] = o.get('ocr_text')
                    elif att['found_status'] != 'found':
                        att['reason'] = 'Referenced in communication but not found in case media folder'
                    attachment_rows.append(att)
                # email attachment rows if parser produced list
                for eatt in row.pop('attachments', []) if isinstance(row.get('attachments', []), list) else []:
                    name = eatt.get('filename') or 'attachment.bin'
                    candidates = media_index.get(os.path.basename(name).lower(), [])
                    saved_path = candidates[0] if isinstance(candidates, list) and candidates else (candidates or '')
                    rel_saved_path = _rel_to_case(case_dir, saved_path) if saved_path else ''
                    attachment_rows.append({
                        'id': hashlib.sha256(f"{row['id']}||{name}".encode()).hexdigest()[:32],
                        'communication_id': row['id'], 'mode': row.get('mode'), 'source_file': row.get('source_file'),
                        'attachment_name': name, 'attachment': name, 'attachment_path': rel_saved_path, 'media_path': rel_saved_path,
                        'found_status': 'found' if saved_path and os.path.exists(saved_path) else 'missing',
                        'preview_image_path': _rel_to_case(case_dir, preview_image_map.get(os.path.basename(name).lower(), '')) if preview_image_map else '',
                        'preview_doc_path': _rel_to_case(case_dir, preview_doc_map.get(os.path.basename(name).lower(), '')) if preview_doc_map else '',
                        'preview_media_path': _rel_to_case(case_dir, preview_media_map.get(os.path.basename(name).lower(), '')) if preview_media_map else '',
                        'attachment_type': '', 'file_ext': os.path.splitext(name)[1].lower(),
                        'ocr_status': '', 'ocr_text':'', 'text_detected':'no', 'reason':'', 'width':None, 'height':None,
                        'gps_lat': None, 'gps_lon': None, 'gps_alt': None, 'gps_timestamp': '', 'gps_source': '', 'gps_confidence': None, 'exif_make':'', 'exif_model':'', 'exif_datetime_original':'', 'has_gps':0,
                    })
                records.append(row)
    # create standalone image/media records for copied files not represented in parsed communications
    represented_paths = set()
    represented_sources = set()
    for a in attachment_rows:
        ap = str(a.get('attachment_path') or a.get('media_path') or '').strip()
        if ap:
            represented_paths.add(ap.replace('\\','/'))
    for r in records:
        sf = str(r.get('source_file') or '').strip()
        if sf:
            represented_sources.add(os.path.normpath(sf))

    if source_index:
        for target_path, src_path in source_index.items():
            ext = os.path.splitext(target_path)[1].lower()
            rel_target = _rel_to_case(case_dir, target_path)
            src_norm = os.path.normpath(src_path)
            if rel_target in represented_paths or src_norm in represented_sources:
                continue
            if ext not in MEDIA_EXTENSIONS:
                continue
            mode = 'photos' if ext in IMAGE_EXTS else 'files'
            if not selected_modes.get(mode, True):
                continue
            base = os.path.basename(target_path)
            geo = extract_attachment_exif(target_path) if ext in IMAGE_EXTS else {
                'gps_lat': None, 'gps_lon': None, 'gps_alt': None, 'gps_timestamp': '', 'gps_source': '', 'gps_confidence': None, 'exif_make':'', 'exif_model':'', 'exif_datetime_original':'', 'has_gps':0,
            }
            row = {
                'mode': mode,
                'source_file': src_path,
                'timestamp': geo.get('gps_timestamp') or '',
                'date_str': geo.get('gps_timestamp') or '',
                'sender': '',
                'receiver': '',
                'direction': '',
                'chat': '',
                'subject': base,
                'message': f'Standalone {mode[:-1].lower() if mode.endswith("s") else mode.lower()} discovered in input folder',
                'body': '',
                'category': 'Personal' if mode == 'photos' else 'Promotional',
                'categories': 'Personal' if mode == 'photos' else 'Promotional',
                'transcription_status': '',
                'transcription_reason': '',
                'attachment_name': base,
                'attachment_path': rel_target,
                'media_path': rel_target,
                'media_exists': True,
                'media_deleted': False,
                'preview_image_path': _rel_to_case(case_dir, preview_image_map.get(os.path.basename(base).lower(), '')) if preview_image_map else '',
                'preview_doc_path': _rel_to_case(case_dir, preview_doc_map.get(os.path.basename(base).lower(), '')) if preview_doc_map else '',
                'preview_media_path': _rel_to_case(case_dir, preview_media_map.get(os.path.basename(base).lower(), '')) if preview_media_map else '',
            }
            row['id'] = normalize_record_id(row)
            records.append(row)
            attachment_rows.append({
                'id': hashlib.sha256(f"{row['id']}||{base}||{target_path}".encode()).hexdigest()[:32],
                'communication_id': row['id'],
                'mode': mode,
                'source_file': src_path,
                'attachment_name': base,
                'attachment': base,
                'attachment_path': rel_target,
                'media_path': rel_target,
                'found_status': 'found',
                'preview_image_path': row.get('preview_image_path',''),
                'preview_doc_path': row.get('preview_doc_path',''),
                'preview_media_path': row.get('preview_media_path',''),
                'attachment_type': ext.lstrip('.'),
                'file_ext': ext,
                'ocr_status': '', 'ocr_text':'', 'text_detected':'no', 'reason':'', 'width':None, 'height':None,
                **geo,
            })
            represented_paths.add(rel_target)
            represented_sources.add(src_norm)

    # dedupe attachments and set attachment_count
    dedup_att=[]; seena=set()
    for a in attachment_rows:
        k=(a['communication_id'],a['attachment_name'],a['attachment_path'])
        if k in seena: continue
        seena.add(k); dedup_att.append(a)
    for a in dedup_att:
        if a.get('attachment_path'):
            try:
                a['attachment_path'] = _rel_to_case(case_dir, a.get('attachment_path')) or a.get('attachment_path')
                a['media_path'] = a.get('attachment_path')
            except Exception:
                pass
    counts={}
    for a in dedup_att: counts[a['communication_id']] = counts.get(a['communication_id'],0)+1
    for r in records: r['attachment_count'] = counts.get(r['id'],0)
    return records, dedup_att


def run_analysis(input_path, output_path, case_no='CASE-001', progress=None, transcribe_audio=False, audio_max_transcription_seconds=None, selected_models=None, selected_modes=None):
    if not input_path or not os.path.exists(input_path):
        raise FileNotFoundError(f'Input path not found: {input_path}')
    case_dir, media_dir, dashboard_dir, logs_dir, extracted_dir, preview_dir = ensure_dirs(output_path, case_no)
    def emit(msg):
        if callable(progress): progress(msg)
        else: print(msg)
    stage_times = {}
    overall_started = time.perf_counter()
    selected_modes = _normalize_selected_modes(selected_modes)
    emit(f"[progress] Starting analysis: input={input_path}")
    emit(f"[progress] Case output: {case_dir}")
    enabled_modes = [k for k,v in selected_modes.items() if v]
    disabled_modes = [k for k,v in selected_modes.items() if not v]
    emit(f"[progress] Enabled modes: {', '.join(enabled_modes)}")
    if disabled_modes:
        emit(f"[progress] Skipped modes: {', '.join(disabled_modes)}")
    selected_models = [str(x).strip().lower() for x in (selected_models or []) if str(x).strip()] or ['generic']
    category_config = load_category_config(os.path.dirname(__file__), os.path.join(case_dir,'config','categories'), selected_models=selected_models)
    case_cfg_dir = os.path.join(case_dir, 'config')
    os.makedirs(case_cfg_dir, exist_ok=True)
    try:
        with open(os.path.join(case_cfg_dir, 'effective_categories.json'), 'w', encoding='utf-8') as f:
            json.dump(category_config, f, indent=2)
    except Exception:
        pass
    t0 = time.perf_counter()
    media_index, copied_media, scan_roots, source_index = build_media_index(input_path, extracted_dir, media_dir, emit, logs_dir=logs_dir)
    stage_times['media_index'] = time.perf_counter() - t0
    emit(f"[progress] Media consolidated: {len(copied_media)} file(s)")
    t0 = time.perf_counter()
    preview_image_map, preview_doc_map, preview_media_map = generate_preview_assets(copied_media, preview_dir, emit)
    stage_times['preview_assets'] = time.perf_counter() - t0
    emit(f"[progress] Preview cache built: images={len(preview_image_map)}, docs={len(preview_doc_map)}, media={len(preview_media_map)}")
    try:
        t0 = time.perf_counter()
        image_infos = []
        for media_path in copied_media:
            if os.path.splitext(media_path)[1].lower() in IMAGE_EXTS:
                image_infos.append(extract_image_info(media_path))
        if image_infos:
            photo_csv = os.path.join(logs_dir, 'photo_exif_metadata.csv')
            export_to_csv(image_infos, photo_csv)
            emit(f"[progress] Photo EXIF metadata CSV written: {photo_csv}")
        stage_times['photo_exif_csv'] = time.perf_counter() - t0
    except Exception as exc:
        emit(f"[progress] Photo EXIF metadata CSV skipped: {exc}")
    parser_timing = {}
    t0 = time.perf_counter()
    records, attachments = collect_records(scan_roots, media_index, category_config, emit, case_dir, transcribe_audio=transcribe_audio, audio_max_transcription_seconds=audio_max_transcription_seconds, preview_image_map=preview_image_map, preview_doc_map=preview_doc_map, preview_media_map=preview_media_map, source_index=source_index, selected_modes=selected_modes, parser_timing=parser_timing)
    stage_times['collect_records'] = time.perf_counter() - t0
    missing=[a for a in attachments if a.get('found_status') != 'found']
    emit(f"[progress] Communication records collected: {len(records)}")
    standalone_count = len([r for r in records if str(r.get('mode')).lower() in ('photos','files')])
    emit(f"[progress] Standalone media/file records added: {standalone_count}")
    emit(f"[progress] Attachments linked: {len(attachments)}")
    emit(f"[progress] Missing artifacts: {len(missing)}")

    db_path = os.path.join(case_dir, 'mxa_comm.db')
    t0 = time.perf_counter()
    db = DB(db_path)
    id_map={}
    for r in records:
        comm_id = db.add_comm(r)
        id_map[r['id']] = comm_id
    for a in attachments:
        aa = dict(a)
        aa['communication_id'] = id_map.get(a['communication_id'])
        db.add_attachment(aa)
    stage_times['sqlite_write'] = time.perf_counter() - t0
    try:
        t0 = time.perf_counter()
        gps_log_path = os.path.join(logs_dir, 'gps_extraction.log')
        updated_meta = db.backfill_image_metadata(progress_cb=emit, log_path=gps_log_path)
        emit(f"[progress] Image EXIF/GPS backfill completed: {updated_meta} attachment(s) updated")
        emit(f"[progress] GPS extraction log written: {gps_log_path}")
        stage_times['gps_backfill'] = time.perf_counter() - t0
    except Exception as exc:
        emit(f"[progress] Image metadata backfill skipped: {exc}")
    db.close()
    emit(f"[progress] SQLite DB written: {db_path}")
    try:
        t0 = time.perf_counter()
        generate_dashboard(records, attachments, missing, case_dir)
        emit("[progress] Dashboard generation complete")
        stage_times['dashboard'] = time.perf_counter() - t0
    except Exception as exc:
        emit(f"[error] Dashboard generation failed: {exc!r}")
    with open(os.path.join(logs_dir,'run.log'),'w',encoding='utf-8') as f:
        f.write('\n'.join([
            f"communications={len(records)}", f"attachments={len(attachments)}", f"missing={len(missing)}"
        ]))
    return {'case_dir': case_dir, 'db_path': db_path, 'records': len(records), 'attachments': len(attachments), 'missing': len(missing), 'media_files': len(copied_media)}